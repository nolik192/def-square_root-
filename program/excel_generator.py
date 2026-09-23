"""Tworzy arkusz excel/metoda_babilonska.xlsx z implementacją metody babilońskiej.

Uruchomienie:  python3 excel_generator.py   (wymaga: pip install openpyxl)
"""
import os
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "excel", "metoda_babilonska.xlsx")
thin = Side(style="thin", color="B7B0A2")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="1F3A5F")
input_fill = PatternFill("solid", fgColor="FCE9D9")
NUM = "0.000000000000000"

wb = Workbook()
ws = wb.active
ws.title = "Metoda babilonska"
ws["A1"] = "Metoda babilońska (Herona): x(n+1) = ( x(n) + S / x(n) ) / 2"
ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
ws["A3"], ws["B3"] = "Liczba S:", 2
ws["A4"], ws["B4"] = "Przybliżenie początkowe x0:", "=B3"
ws["A5"], ws["B5"] = "Dokładność (eps):", 1e-15
ws["B5"].number_format = "0.0E+00"
for c in ("B3", "B4", "B5"):
    ws[c].fill = input_fill
    ws[c].border = box
ws["D3"], ws["E3"] = "Funkcja PIERWIASTEK():", "=SQRT(B3)"
ws["D4"], ws["E4"] = "Liczba iteracji:", '=COUNTIF(G10:G28,"nie")+1'
ws["E3"].number_format = NUM
for c in ("A3", "A4", "A5", "D3", "D4"):
    ws[c].font = Font(bold=True)

headers = ["n", "x(n)", "x(n)^2", "S / x(n)", "|x(n) - x(n-1)|", "Błąd |x(n) - √S|", "Koniec?"]
for i, h in enumerate(headers, start=1):
    c = ws.cell(row=8, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = head_fill
    c.alignment = Alignment(horizontal="center")
    c.border = box
for k in range(20):
    r = 9 + k
    ws.cell(row=r, column=1, value=k)
    ws.cell(row=r, column=2, value="=B4" if k == 0 else f"=(B{r-1}+$B$3/B{r-1})/2")
    ws.cell(row=r, column=3, value=f"=B{r}^2")
    ws.cell(row=r, column=4, value=f"=$B$3/B{r}")
    ws.cell(row=r, column=5, value="" if k == 0 else f"=ABS(B{r}-B{r-1})")
    ws.cell(row=r, column=6, value=f"=ABS(B{r}-SQRT($B$3))")
    ws.cell(row=r, column=7, value="" if k == 0 else f'=IF(E{r}<=$B$5*B{r},"tak","nie")')
    for col in range(1, 8):
        cell = ws.cell(row=r, column=col)
        cell.border = box
        if 2 <= col <= 4:
            cell.number_format = NUM
        elif col in (5, 6):
            cell.number_format = "0.00E+00"
ws.column_dimensions["A"].width = 28
for col in "BCDEF":
    ws.column_dimensions[col].width = 24
ws.column_dimensions["G"].width = 10
ws.freeze_panes = "A9"

ch = LineChart()
ch.title = "Kolejne przybliżenia x(n)"
ch.y_axis.title = "x(n)"
ch.x_axis.title = "n"
ch.add_data(Reference(ws, min_col=2, min_row=8, max_row=18), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=9, max_row=18))
ch.height, ch.width = 8, 16
ws.add_chart(ch, "I3")

# Arkusz 2: pierwiastki kilku liczb naraz
ws2 = wb.create_sheet("Wiele liczb")
ws2["A1"] = "Pierwiastki kilku liczb – każda kolumna to osobna liczba S"
ws2["A1"].font = Font(bold=True, size=14, color="1F3A5F")
ws2["A3"] = "S ="
ws2["A4"] = "n  \\  x(n)"
for c in ("A3", "A4"):
    ws2[c].font = Font(bold=True)
for j, s in enumerate([2, 3, 10, 100, 12345, 1000000]):
    col = 2 + j
    L = ws2.cell(row=3, column=col).column_letter
    ws2.cell(row=3, column=col, value=s).fill = input_fill
    for k in range(18):
        r = 5 + k
        ws2.cell(row=r, column=1, value=k)
        f = f"={L}3" if k == 0 else f"=({L}{r-1}+{L}$3/{L}{r-1})/2"
        ws2.cell(row=r, column=col, value=f).number_format = NUM
    ws2.cell(row=24, column=col, value=f"=SQRT({L}3)").number_format = NUM
    ws2.column_dimensions[L].width = 26
ws2["A24"] = "PIERWIASTEK()"
ws2["A24"].font = Font(bold=True)
ws2.column_dimensions["A"].width = 16

wb.save(OUT)
print("Zapisano", os.path.abspath(OUT))
